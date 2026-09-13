from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path

from flask import Flask, flash, redirect, render_template, request, send_file, url_for
from openpyxl import load_workbook
from werkzeug.utils import secure_filename

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
MODEL_PATH = BASE_DIR / "modelo_lab217.xlsx"

app = Flask(__name__)
app.secret_key = "lab217-planilhas"

SHEETS = {
    "Controle": {
        "key_col": 2,
        "fields": [
            ("reagente", "Reagente", "text", 2),
            ("data", "Data", "date", 3),
            ("para_quem", "Para quem", "text", 4),
            ("observacao", "Observação", "textarea", 5),
            ("preco", "Preço (R$)", "money", 6),
            ("forma_compra", "Forma de compra", "select:À vista|Parcelado|FAPEMIG - reembolso integral|FAPEMIG - pagamento direto|FAPEMIG - parcial|Outro", 10),
        ],
    },
    "Pagamentos": {
        "key_col": 2,
        "fields": [
            ("id_item", "ID do item", "number", 2),
            ("data_pagamento", "Data do pagamento", "date", 3),
            ("valor_pago", "Valor pago (R$)", "money", 4),
            ("observacao", "Observação", "textarea", 5),
        ],
    },
    "Observações": {
        "key_col": 4,
        "fields": [
            ("data", "Data", "date", 2),
            ("id_item", "ID do item (opcional)", "number", 3),
            ("observacao", "Observação", "textarea", 4),
        ],
    },
}


def ensure_data_dir():
    DATA_DIR.mkdir(exist_ok=True)
    if not any(DATA_DIR.glob("*.xlsx")) and MODEL_PATH.exists():
        shutil.copy2(MODEL_PATH, DATA_DIR / "Controle_Lab217.xlsx")


def safe_name(name: str) -> str:
    name = secure_filename((name or "").strip()) or "Controle_Lab217"
    if not name.lower().endswith(".xlsx"):
        name += ".xlsx"
    return name


def path_for(filename: str) -> Path:
    path = (DATA_DIR / safe_name(filename)).resolve()
    if DATA_DIR.resolve() not in path.parents:
        raise ValueError("Nome inválido")
    return path


def parse_value(kind: str, raw: str):
    raw = (raw or "").strip()
    if raw == "":
        return None
    if kind in ("number", "money"):
        txt = raw.replace("R$", "").replace(" ", "")
        if "," in txt:
            txt = txt.replace(".", "").replace(",", ".")
        return float(txt)
    if kind == "date":
        return datetime.strptime(raw, "%Y-%m-%d").date()
    return raw


def set_calc_on_load(wb):
    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        wb.calculation.calcMode = "auto"
    except Exception:
        pass


def first_free_row(ws, key_col: int) -> int:
    # IDs são pré-reservados; procuramos a primeira linha cujo conteúdo principal está vazio.
    for row in range(2, ws.max_row + 1):
        if ws.cell(row, key_col).value in (None, ""):
            return row
    # Se acabarem as linhas reservadas, cria uma nova com ID sequencial.
    row = ws.max_row + 1
    ids = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)]
    ids = [int(v) for v in ids if isinstance(v, (int, float))]
    ws.cell(row, 1).value = max(ids, default=0) + 1
    return row


def ensure_control_formulas(ws, row: int):
    ws.cell(row, 7).value = f'=IF(B{row}="","",SUMIF(Pagamentos!$B$2:$B$1001,A{row},Pagamentos!$D$2:$D$1001))'
    ws.cell(row, 8).value = f'=IF(B{row}="","",MAX(F{row}-G{row},0))'
    ws.cell(row, 9).value = f'=IF(B{row}="","",IF(H{row}<=0,"PAGO",IF(G{row}>0,"PARCIAL","PENDENTE")))'
    for c in (6, 7, 8):
        ws.cell(row, c).number_format = 'R$ #,##0.00'


def save(wb, path: Path):
    set_calc_on_load(wb)
    wb.save(path)


def fmt_money(v):
    try:
        return f"R$ {float(v):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    except (TypeError, ValueError):
        return "R$ 0,00"


def fmt_date(v):
    if hasattr(v, "strftime"):
        return v.strftime("%d/%m/%Y")
    return str(v or "")


def get_payments(wb):
    totals = {}
    if "Pagamentos" not in wb.sheetnames:
        return totals
    ws = wb["Pagamentos"]
    for r in range(2, ws.max_row + 1):
        item_id = ws.cell(r, 2).value
        val = ws.cell(r, 4).value
        if item_id in (None, "") or val in (None, ""):
            continue
        try:
            item_id = int(float(item_id))
            val = float(val)
        except (TypeError, ValueError):
            continue
        totals[item_id] = totals.get(item_id, 0.0) + val
    return totals


def rows_for(path: Path, sheet_name: str):
    wb = load_workbook(path, data_only=False)
    payments = get_payments(wb)
    ws = wb[sheet_name]
    rows = []

    if sheet_name == "Controle":
        for r in range(2, ws.max_row + 1):
            reagente = ws.cell(r, 2).value
            if reagente in (None, ""):
                continue
            item_id = int(ws.cell(r, 1).value)
            preco = float(ws.cell(r, 6).value or 0)
            pago = payments.get(item_id, 0.0)
            saldo = max(preco - pago, 0)
            status = "PAGO" if saldo <= 0 else ("PARCIAL" if pago > 0 else "PENDENTE")
            rows.append({"excel_row": r, "id": item_id, "values": [
                item_id,
                reagente,
                fmt_date(ws.cell(r, 3).value),
                ws.cell(r, 4).value or "",
                ws.cell(r, 5).value or "",
                fmt_money(preco),
                fmt_money(pago),
                fmt_money(saldo),
                status,
                ws.cell(r, 10).value or "",
            ], "status": status})

        total = sum(float(ws.cell(r, 6).value or 0) for r in range(2, ws.max_row + 1) if ws.cell(r, 2).value not in (None, ""))
        pago_total = sum(payments.get(int(ws.cell(r, 1).value), 0.0) for r in range(2, ws.max_row + 1) if ws.cell(r, 2).value not in (None, ""))
        divida = sum(max(float(ws.cell(r, 6).value or 0) - payments.get(int(ws.cell(r, 1).value), 0.0), 0) for r in range(2, ws.max_row + 1) if ws.cell(r, 2).value not in (None, ""))
        totals = {"total": total, "pago": pago_total, "divida": divida}
        headers = ["ID", "Reagente", "Data", "Para quem", "Observação", "Preço", "Total pago", "Saldo em aberto", "Status", "Forma de compra"]

    elif sheet_name == "Pagamentos":
        for r in range(2, ws.max_row + 1):
            if ws.cell(r, 2).value in (None, ""):
                continue
            rows.append({"excel_row": r, "id": int(ws.cell(r, 1).value), "values": [
                int(ws.cell(r, 1).value),
                ws.cell(r, 2).value,
                fmt_date(ws.cell(r, 3).value),
                fmt_money(ws.cell(r, 4).value or 0),
                ws.cell(r, 5).value or "",
            ]})
        totals = None
        headers = ["ID Pagamento", "ID do item", "Data do pagamento", "Valor pago", "Observação"]

    elif sheet_name == "Observações":
        for r in range(2, ws.max_row + 1):
            if ws.cell(r, 4).value in (None, ""):
                continue
            rows.append({"excel_row": r, "id": int(ws.cell(r, 1).value), "values": [
                int(ws.cell(r, 1).value),
                fmt_date(ws.cell(r, 2).value),
                ws.cell(r, 3).value or "",
                ws.cell(r, 4).value or "",
            ]})
        totals = None
        headers = ["ID", "Data", "ID do item", "Observação"]
    else:
        totals = None
        headers = []
    return wb, rows, totals, headers


@app.route("/")
def index():
    ensure_data_dir()
    files = sorted(DATA_DIR.glob("*.xlsx"), key=lambda p: p.stat().st_mtime, reverse=True)
    return render_template("index.html", planilhas=files)


@app.post("/nova")
def nova_planilha():
    ensure_data_dir()
    destino = path_for(request.form.get("nome", ""))
    if destino.exists():
        flash("Já existe uma planilha com esse nome.", "erro")
        return redirect(url_for("index"))
    shutil.copy2(MODEL_PATH, destino)
    flash("Planilha criada a partir do modelo inteligente do Lab 217.", "sucesso")
    return redirect(url_for("ver_planilha", filename=destino.name, aba="Controle"))


@app.post("/upload")
def upload_planilha():
    ensure_data_dir()
    arq = request.files.get("arquivo")
    if not arq or not arq.filename or not arq.filename.lower().endswith(".xlsx"):
        flash("Selecione um arquivo .xlsx.", "erro")
        return redirect(url_for("index"))
    destino = path_for(arq.filename)
    arq.save(destino)
    try:
        wb = load_workbook(destino, data_only=False)
        obrigatorias = {"Controle", "Pagamentos", "Resumo", "Observações"}
        if not obrigatorias.issubset(set(wb.sheetnames)):
            raise ValueError("modelo incompatível")
        save(wb, destino)
    except Exception:
        destino.unlink(missing_ok=True)
        flash("Essa planilha não possui as 4 abas esperadas do modelo do Lab 217.", "erro")
        return redirect(url_for("index"))
    flash("Planilha enviada e salva.", "sucesso")
    return redirect(url_for("ver_planilha", filename=destino.name, aba="Controle"))


@app.route("/planilha/<filename>")
def ver_planilha(filename):
    path = path_for(filename)
    if not path.exists():
        flash("Planilha não encontrada.", "erro")
        return redirect(url_for("index"))
    aba = request.args.get("aba", "Controle")
    wb = load_workbook(path, data_only=False)
    if aba not in wb.sheetnames:
        aba = "Controle"
    if aba == "Resumo":
        payments = get_payments(wb)
        ws = wb["Controle"]
        items = [(int(ws.cell(r, 1).value), float(ws.cell(r, 6).value or 0)) for r in range(2, ws.max_row + 1) if ws.cell(r, 2).value not in (None, "")]
        total = sum(v for _, v in items)
        pago = sum(payments.get(i, 0) for i, _ in items)
        divida = sum(max(v - payments.get(i, 0), 0) for i, v in items)
        return render_template("resumo.html", filename=path.name, abas=wb.sheetnames, total=total, pago=pago, divida=divida)

    _wb, rows, totals, headers = rows_for(path, aba)
    editable = aba in SHEETS
    return render_template("planilha.html", filename=path.name, abas=wb.sheetnames, aba=aba,
                           rows=rows, totals=totals, headers=headers, editable=editable,
                           fields=SHEETS.get(aba, {}).get("fields", []))


@app.post("/planilha/<filename>/<aba>/adicionar")
def adicionar(filename, aba):
    path = path_for(filename)
    if aba not in SHEETS or not path.exists():
        return redirect(url_for("index"))
    wb = load_workbook(path, data_only=False)
    ws = wb[aba]
    config = SHEETS[aba]
    row = first_free_row(ws, config["key_col"])
    for name, label, kind, col in config["fields"]:
        try:
            value = parse_value(kind, request.form.get(name, ""))
        except ValueError:
            flash(f"Valor inválido em {label}.", "erro")
            return redirect(url_for("ver_planilha", filename=filename, aba=aba))
        ws.cell(row, col).value = value
        if kind == "date" and value:
            ws.cell(row, col).number_format = "dd/mm/yyyy"
        if kind == "money":
            ws.cell(row, col).number_format = 'R$ #,##0.00'
    if aba == "Controle":
        ensure_control_formulas(ws, row)
    save(wb, path)
    flash("Registro adicionado diretamente à planilha.", "sucesso")
    return redirect(url_for("ver_planilha", filename=filename, aba=aba))


@app.route("/planilha/<filename>/<aba>/<int:row>/editar", methods=["GET", "POST"])
def editar(filename, aba, row):
    path = path_for(filename)
    if aba not in SHEETS or not path.exists():
        return redirect(url_for("index"))
    wb = load_workbook(path, data_only=False)
    ws = wb[aba]
    fields = SHEETS[aba]["fields"]
    if request.method == "POST":
        for name, label, kind, col in fields:
            try:
                value = parse_value(kind, request.form.get(name, ""))
            except ValueError:
                flash(f"Valor inválido em {label}.", "erro")
                return redirect(request.url)
            ws.cell(row, col).value = value
            if kind == "date" and value:
                ws.cell(row, col).number_format = "dd/mm/yyyy"
            if kind == "money":
                ws.cell(row, col).number_format = 'R$ #,##0.00'
        if aba == "Controle":
            ensure_control_formulas(ws, row)
        save(wb, path)
        flash("Registro alterado.", "sucesso")
        return redirect(url_for("ver_planilha", filename=filename, aba=aba))
    values = {}
    for name, _label, kind, col in fields:
        v = ws.cell(row, col).value
        if kind == "date" and hasattr(v, "strftime"):
            v = v.strftime("%Y-%m-%d")
        values[name] = "" if v is None else v
    return render_template("editar.html", filename=filename, aba=aba, row=row, fields=fields, values=values, registro_id=ws.cell(row,1).value)


@app.post("/planilha/<filename>/<aba>/<int:row>/excluir")
def excluir(filename, aba, row):
    path = path_for(filename)
    if aba not in SHEETS or not path.exists():
        return redirect(url_for("index"))
    wb = load_workbook(path, data_only=False)
    ws = wb[aba]
    record_id = ws.cell(row, 1).value
    if aba == "Controle":
        # Limpa o objeto, preserva seu ID, e remove pagamentos vinculados para não deixar valores órfãos.
        for c in range(2, 11):
            ws.cell(row, c).value = None
        ensure_control_formulas(ws, row)
        pws = wb["Pagamentos"]
        for r in range(2, pws.max_row + 1):
            try:
                linked = int(float(pws.cell(r, 2).value)) == int(record_id)
            except (TypeError, ValueError):
                linked = False
            if linked:
                for c in range(2, 6):
                    pws.cell(r, c).value = None
    else:
        max_col = 5 if aba == "Pagamentos" else 4
        for c in range(2, max_col + 1):
            ws.cell(row, c).value = None
    save(wb, path)
    flash("Registro excluído; o ID reservado foi preservado.", "sucesso")
    return redirect(url_for("ver_planilha", filename=filename, aba=aba))


@app.get("/planilha/<filename>/baixar")
def baixar(filename):
    path = path_for(filename)
    return send_file(path, as_attachment=True, download_name=path.name)


@app.post("/planilha/<filename>/duplicar")
def duplicar(filename):
    origem = path_for(filename)
    destino = path_for(request.form.get("novo_nome", ""))
    if destino.exists():
        flash("Já existe uma planilha com esse nome.", "erro")
    else:
        shutil.copy2(origem, destino)
        flash("Cópia criada.", "sucesso")
    return redirect(url_for("index"))


@app.post("/planilha/<filename>/excluir-arquivo")
def excluir_arquivo(filename):
    path = path_for(filename)
    path.unlink(missing_ok=True)
    flash("Arquivo excluído.", "sucesso")
    return redirect(url_for("index"))


if __name__ == "__main__":
    ensure_data_dir()
    app.run(debug=True)
